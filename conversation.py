import os
import datetime
from docx import Document
from fpdf import FPDF

class ConversationManager:
    def __init__(self, filename="conversation_history.txt"):
        self.filename = filename
        self.history = []
        self.file_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), self.filename)
        # Ensure the file exists
        self._init_history_file()

    def _init_history_file(self):
        """Initializes the text file with a header if it doesn't exist."""
        if not os.path.exists(self.file_path):
            with open(self.file_path, "w", encoding="utf-8") as f:
                f.write(f"=== Voice Assistant Chat History (Started: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}) ===\n\n")

    def add_message(self, role: str, content: str):
        """Adds a message to the history and appends to the text file."""
        timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        self.history.append({
            "role": role,
            "content": content,
            "timestamp": timestamp
        })
        self._append_to_file(role, content, timestamp)

    def _append_to_file(self, role: str, content: str, timestamp: str):
        """Appends a single chat message to the history file."""
        display_role = "User" if role == "user" else "Assistant"
        try:
            with open(self.file_path, "a", encoding="utf-8") as f:
                f.write(f"[{timestamp}] {display_role}:\n{content}\n\n")
        except Exception as e:
            print(f"Failed to append to history file: {e}")

    def clear(self):
        """Clears the session history and overwrites the text log file."""
        self.history = []
        try:
            with open(self.file_path, "w", encoding="utf-8") as f:
                f.write(f"=== Voice Assistant Chat History (Cleared & Restarted: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}) ===\n\n")
        except Exception as e:
            print(f"Failed to reset history file: {e}")

    def get_history(self) -> list:
        """Returns the in-memory conversation history."""
        return self.history

    def get_formatted_history(self) -> list:
        """Returns the history formatted for LLM context (role and content)."""
        return [{"role": msg["role"], "content": msg["content"]} for msg in self.history]

    def _sanitize_for_latin1(self, text: str) -> str:
        """Sanitizes text to avoid Latin-1 encoding issues with standard FPDF fonts."""
        # Replace common unicode quotation marks and dashes
        replacements = {
            '\u201c': '"', '\u201d': '"',
            '\u2018': "'", '\u2019': "'",
            '\u2013': '-', '\u2014': '-',
            '\u2212': '-', '\u2022': '*',
            '\u2026': '...'
        }
        for uni_char, ascii_char in replacements.items():
            text = text.replace(uni_char, ascii_char)
        # Encode to latin-1 with replacement for characters we missed
        return text.encode('latin-1', 'replace').decode('latin-1')

    def export_to_pdf(self, export_path: str):
        """Exports the chat history to a formatted PDF file."""
        try:
            pdf = FPDF()
            pdf.add_page()
            
            # Title
            pdf.set_font("Helvetica", style="B", size=16)
            pdf.cell(0, 10, "Voice Assistant Chat History", ln=True, align="C")
            pdf.set_font("Helvetica", size=10)
            pdf.cell(0, 10, f"Exported on: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", ln=True, align="C")
            pdf.line(10, 30, 200, 30)
            pdf.ln(10)
            
            for msg in self.history:
                role_label = "User" if msg["role"] == "user" else "Assistant"
                timestamp = msg.get("timestamp", "")
                
                # Format role header
                pdf.set_font("Helvetica", style="B", size=11)
                if role_label == "User":
                    pdf.set_text_color(0, 102, 204)  # Blue for User
                else:
                    pdf.set_text_color(46, 125, 50)  # Green for Assistant
                    
                pdf.cell(0, 6, f"{role_label} ({timestamp})", ln=True)
                
                # Format message body
                pdf.set_font("Helvetica", size=10)
                pdf.set_text_color(33, 33, 33)  # Dark Gray
                
                cleaned_content = self._sanitize_for_latin1(msg["content"])
                pdf.multi_cell(0, 6, cleaned_content)
                pdf.ln(4)
                
            pdf.output(export_path)
            return True
        except Exception as e:
            print(f"Error exporting PDF: {e}")
            raise e

    def export_to_docx(self, export_path: str):
        """Exports the chat history to a Microsoft Word document."""
        try:
            doc = Document()
            
            # Title
            doc.add_heading("Voice Assistant Chat History", level=0)
            doc.add_paragraph(f"Exported on: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
            
            for msg in self.history:
                role_label = "User" if msg["role"] == "user" else "Assistant"
                timestamp = msg.get("timestamp", "")
                
                # Add Paragraph
                p = doc.add_paragraph()
                
                # Add bold role and timestamp
                r_role = p.add_run(f"[{timestamp}] {role_label}: ")
                r_role.bold = True
                if role_label == "User":
                    r_role.font.color.rgb = None  # standard
                
                # Add content
                p.add_run(f"\n{msg['content']}")
                
            doc.save(export_path)
            return True
        except Exception as e:
            print(f"Error exporting Word document: {e}")
            raise e
